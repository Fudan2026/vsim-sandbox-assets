# -*- coding: utf-8 -*-
"""教学指导配套 PPT v2 —— 视觉体系蒸馏自主人讲座模板：
白底卡片 + 顶部彩色细条 + 圆角 + 右侧页码，强调色轮换（蓝/橙/翠绿/玫红/紫），
正文 Hiragino Sans GB，封面衬线点缀 Hoefler Text，字号层级 9–15/22/30/100。
内容口径与 v1 完全一致。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- 设计令牌（讲座模板色板）----
INK    = RGBColor(0x0F, 0x17, 0x2A)   # slate-900 标题
GRAY   = RGBColor(0x33, 0x41, 0x55)   # slate-700 正文
FAINT  = RGBColor(0x64, 0x74, 0x8B)   # slate-500 弱化
BORDER = RGBColor(0xE2, 0xE8, 0xF0)   # slate-200 卡片边
TINT   = RGBColor(0xF1, 0xF5, 0xF9)   # slate-100 浅底
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

BLUE   = RGBColor(0x25, 0x63, 0xEB)   # 主蓝
ORANGE = RGBColor(0xF9, 0x73, 0x16)
EMER   = RGBColor(0x10, 0xB9, 0x81)
ROSE   = RGBColor(0xF4, 0x3F, 0x5E)
VIOLET = RGBColor(0x7C, 0x3A, 0xED)
AMBER  = RGBColor(0xF5, 0x9E, 0x0B)
CYAN   = RGBColor(0x06, 0xB6, 0xD4)

BLUE_T  = RGBColor(0xEF, 0xF6, 0xFF)  # blue-50
ORAN_T  = RGBColor(0xFF, 0xF7, 0xED)  # orange-50
EMER_T  = RGBColor(0xEC, 0xFD, 0xF5)  # emerald-50
ROSE_T  = RGBColor(0xFF, 0xF1, 0xF3)  # rose-50
VIOL_T  = RGBColor(0xF5, 0xF3, 0xFF)  # violet-50
AMBER_T = RGBColor(0xFF, 0xFB, 0xEB)  # amber-50
SKY     = RGBColor(0x38, 0xBD, 0xF8)  # sky-400 装饰

FONT  = "Hiragino Sans GB"
SERIF = "Hoefler Text"
TOTAL = 20

SW, SH = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def _set_font(run, size, color, bold=False, font=FONT):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, font
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', font)

def box(s, x, y, w, h, fill=None, line=None, round_=False, line_w=1.0, adj=0.09):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if round_:
        try: shp.adjustments[0] = adj
        except Exception: pass
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp

def oval(s, x, y, d, fill):
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.fill.background(); shp.shadow.inherit = False
    return shp

def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0, space_after=2):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        for run_ in para:
            t, sz, c, b = run_[:4]
            fnt = run_[4] if len(run_) > 4 else FONT
            r = p.add_run(); r.text = t
            _set_font(r, sz, c, b, fnt)
    return tb

def header(s, kicker, title, accent, sub=None):
    """讲座式页眉：顶部彩条 + 竖条 + 眉题 + 标题(+副题) + 页脚页码"""
    box(s, 0, 0, SW, Inches(0.1), fill=accent)
    box(s, Inches(0.6), Inches(0.5), Inches(0.1), Inches(0.9), fill=accent)
    txt(s, Inches(0.9), Inches(0.4), Inches(9.8), Inches(0.3), [[(kicker, 12, FAINT, False)]])
    txt(s, Inches(0.9), Inches(0.75), Inches(11.6), Inches(0.6), [[(title, 22, INK, True)]])
    if sub:
        txt(s, Inches(0.9), Inches(1.38), Inches(11.9), Inches(0.35), [[(sub, 11.5, FAINT, False)]])

def footer(s, n, accent):
    txt(s, Inches(0.6), Inches(7.08), Inches(8), Inches(0.3),
        [[("AI 选股信号该配多大仓位？ · 虚拟仿真实验教学案例 · 教学指导配套", 8.5, FAINT, False)]])
    txt(s, Inches(11.6), Inches(7.02), Inches(1.2), Inches(0.35),
        [[(f"{n:02d} / {TOTAL}", 10, accent, True)]], align=PP_ALIGN.RIGHT)

def num_badge(s, x, y, n, color, d=0.52):
    oval(s, x, y, Inches(d), color)
    txt(s, x, y + Inches(0.05), Inches(d), Inches(d - 0.08),
        [[(str(n), 15, WHITE, True)]], align=PP_ALIGN.CENTER)

def card(s, x, y, w, h, accent, head=None, body=None, head_size=13, body_size=10.5, strip=True, fill=WHITE):
    box(s, x, y, w, h, fill=fill, line=BORDER, round_=True)
    if strip:
        box(s, x, y, w, Inches(0.09), fill=accent, round_=False)
    pad = Inches(0.22)
    ty = y + Inches(0.18)
    if head:
        txt(s, x + pad, ty, w - pad * 2, Inches(0.4), [[(head, head_size, INK, True)]])
        ty += Inches(0.48)
    if body:
        paras = [[(ln, body_size, GRAY, False)] for ln in body.split("\n")] if "\n" in body else [[(body, body_size, GRAY, False)]]
        txt(s, x + pad, ty, w - pad * 2, y + h - ty - Inches(0.15), paras, spacing=1.14)

def insight(s, formula, note, accent, tint, y=Inches(6.15), h=Inches(0.72)):
    """底部洞察条：公式式断言 + 右侧注解"""
    box(s, Inches(0.6), y, Inches(12.13), h, fill=tint, round_=True)
    txt(s, Inches(0.9), y + Inches(0.08), Inches(8.6), h - Inches(0.14),
        [[(formula, 12.5, INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(9.3), y + Inches(0.08), Inches(3.2), h - Inches(0.14),
        [[(note, 10.5, accent, True)]], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

def divider(sec_no, sec_en, sec_cn, points, accent, tint):
    s = slide()
    box(s, 0, 0, SW, SH, fill=WHITE)
    box(s, 0, 0, SW, Inches(0.1), fill=accent)
    oval(s, Inches(9.6), Inches(-2.4), Inches(6.2), tint)
    oval(s, Inches(11.2), Inches(-1.1), Inches(3.2), accent)
    oval(s, Inches(-1.6), Inches(5.6), Inches(3.8), tint)
    txt(s, Inches(0.9), Inches(1.7), Inches(4), Inches(1.9), [[(sec_no, 96, tint if tint != WHITE else accent, True)]])
    txt(s, Inches(0.95), Inches(3.6), Inches(10), Inches(0.4), [[(sec_en, 14, accent, True)]])
    txt(s, Inches(0.95), Inches(4.05), Inches(10), Inches(0.8), [[(sec_cn, 32, INK, True)]])
    yy = Inches(5.2)
    for p in points:
        txt(s, Inches(0.95), yy, Inches(11.5), Inches(0.35), [[("·  " + p, 12, GRAY, False)]])
        yy += Inches(0.42)
    footer(s, prs_slides_len() + 1, accent)
    return s

def prs_slides_len():
    return len(prs.slides._sldIdLst)

def styled_table(s, x, y, w, headers, rows, col_w=None, font_size=10.5, row_h=0.34, head_fill=INK):
    n_r, n_c = len(rows) + 1, len(headers)
    gt = s.shapes.add_table(n_r, n_c, x, y, w, Inches(row_h * n_r)).table
    if col_w:
        for i, cw in enumerate(col_w): gt.columns[i].width = Inches(cw)
    for j, htxt in enumerate(headers):
        c = gt.cell(0, j)
        c.fill.solid(); c.fill.fore_color.rgb = head_fill
        c.margin_top = c.margin_bottom = Pt(1)
        p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = htxt; _set_font(r, font_size, WHITE, True)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            c = gt.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = WHITE if i % 2 else TINT
            c.margin_top = c.margin_bottom = Pt(1)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if j else PP_ALIGN.LEFT
            r = p.add_run(); r.text = str(val)
            _set_font(r, font_size, GRAY if j else INK, j == 0)
    return gt

def pill(s, x, y, w, text, fill):
    box(s, x, y, Inches(w), Inches(0.42), fill=fill, round_=True, adj=0.5)
    txt(s, x, y + Inches(0.05), Inches(w), Inches(0.32), [[(text, 11, WHITE, True)]], align=PP_ALIGN.CENTER)

# ================= S1 封面 =================
s = slide()
box(s, 0, 0, SW, SH, fill=WHITE)
oval(s, Inches(9.4), Inches(-2.2), Inches(6.0), BLUE_T)
oval(s, Inches(10.8), Inches(-0.9), Inches(3.2), SKY)
oval(s, Inches(-1.4), Inches(5.4), Inches(3.6), ORAN_T)
box(s, Inches(0.6), Inches(0.6), Inches(0.9), Inches(0.09), fill=BLUE)
txt(s, Inches(0.9), Inches(1.05), Inches(11.5), Inches(0.35),
    [[("智享杯 · 全国大学生经济管理类实验案例分析大赛", 13, FAINT, True)]])
txt(s, Inches(0.85), Inches(1.75), Inches(12.2), Inches(1.1),
    [[("AI 选股信号该配多大仓位？", 40, INK, True)]])
txt(s, Inches(0.9), Inches(2.85), Inches(11.8), Inches(0.55),
    [[("——虚拟仿真下的风险预算干预实验", 20, BLUE, True)]])
txt(s, Inches(0.9), Inches(3.5), Inches(11.8), Inches(0.35),
    [[("How Much Should an AI Signal Buy? A Virtual-Simulation Experiment on Risk-Budget Intervention", 11, FAINT, False, SERIF)]])
for i, (t, f) in enumerate([("实验教学案例", BLUE), ("数据分析案例", ORANGE), ("在线 Notebook", EMER), ("一键复现", VIOLET)]):
    pill(s, Inches(0.9 + i * 1.95), Inches(4.35), 1.75, t, f)
info = [("适用课程", "投资学 / 量化投资 / 金融风险管理"),
        ("实验形态", "平台在线 Notebook 虚拟仿真实验"),
        ("配套材料", "教学指导文档 · 代码与数据（种子固定）")]
for i, (k, v) in enumerate(info):
    x = Inches(0.9 + i * 4.0)
    txt(s, x, Inches(5.25), Inches(3.8), Inches(0.3), [[(k, 10.5, FAINT, False)]])
    txt(s, x, Inches(5.58), Inches(3.8), Inches(0.4), [[(v, 12, INK, True)]])
txt(s, Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.35),
    [[("教学指导配套 PPT", 12, GRAY, False)]])
box(s, 0, 0, SW, Inches(0.1), fill=BLUE)
box(s, 0, Inches(7.4), SW, Inches(0.1), fill=EMER)
txt(s, Inches(0.6), Inches(7.06), Inches(8), Inches(0.3), [[("虚拟仿真实验 · 教学演示，不构成投资建议", 8.5, FAINT, False)]])
txt(s, Inches(11.6), Inches(7.0), Inches(1.2), Inches(0.35), [[(f"01 / {TOTAL}", 10, BLUE, True)]], align=PP_ALIGN.RIGHT)

# ================= S2 目录 =================
s = slide()
box(s, 0, 0, SW, SH, fill=WHITE)
box(s, 0, 0, SW, Inches(0.1), fill=BLUE)
box(s, Inches(0.6), Inches(0.5), Inches(0.1), Inches(0.9), fill=BLUE)
txt(s, Inches(0.9), Inches(0.4), Inches(9), Inches(0.3), [[("CONTENTS", 12, FAINT, True)]])
txt(s, Inches(0.9), Inches(0.72), Inches(9), Inches(0.6), [[("目录 · 四幕结构", 22, INK, True)]])
txt(s, Inches(0.9), Inches(1.45), Inches(11.8), Inches(0.35),
    [[("一份把「AI 投研信号」与「仓位决策训练」融合起来的实验教学案例", 12, GRAY, False)]])
toc = [("01", "背景与意义", "三个真问题 · 为什么必须虚拟仿真", BLUE, BLUE_T),
       ("02", "实验设计", "真实行情沙盘 · 透明信号 · 三组仓位规则", EMER, EMER_T),
       ("03", "结果与发现", "成本之税 · 校准失灵 · 保险定价 · 干预线", ORANGE, ORAN_T),
       ("04", "教学组织与评价", "分层任务 · 典型错误引导 · 四维量规", VIOLET, VIOL_T)]
for i, (n, t, d, c, tint) in enumerate(toc):
    x = Inches(0.6 + (i % 2) * 6.25); y = Inches(2.1 + (i // 2) * 2.25)
    box(s, x, y, Inches(5.95), Inches(1.95), fill=WHITE, line=BORDER, round_=True)
    box(s, x, y, Inches(5.95), Inches(0.09), fill=c)
    txt(s, x + Inches(0.3), y + Inches(0.35), Inches(1.4), Inches(1.1), [[(n, 36, c, True)]])
    txt(s, x + Inches(1.75), y + Inches(0.35), Inches(4.0), Inches(0.45), [[(t, 16, INK, True)]])
    txt(s, x + Inches(1.75), y + Inches(0.95), Inches(4.0), Inches(0.8), [[(d, 10.5, GRAY, False)]], spacing=1.15)
footer(s, 2, BLUE)

# ================= S3 章节页 01 =================
divider("01", "BACKGROUND & SIGNIFICANCE", "背景与意义",
        ["三个真问题：AI 说了买，该买多少？亏了谁来管？",
         "预测力 ≠ 盈利能力：分水岭在仓位规则与交易成本",
         "真实公开行情 + 透明模拟信号 = 可控而复杂的实验沙盘"],
        BLUE, BLUE_T)

# ================= S4 三个真问题 =================
s = slide()
box(s, 0, 0, SW, SH, fill=WHITE)
header(s, "01 · 背景与意义", "三个真问题：AI 投研链条上被忽视的「最后一厘米」", BLUE,
       sub="大模型投研工作台批量进入券商与基金，但信号到资金之间的转化环节没有训练环境")
qs = [("来自行业的真实提问 01", "AI 说「买」，仓位该给多少？", "每日数百条看多/看空观点与置信度，信号到仓位的转化全靠交易员手感，缺乏可训练、可考核的实验环境。", BLUE, BLUE_T),
      ("来自行业的真实提问 02", "高置信度信号，值得重仓吗？", "未经校准的置信度直接映射为仓位，会把模型的过度自信放大为账户的过度敞口——谁来验证置信度的成色？", ROSE, ROSE_T),
      ("来自行业的真实提问 03", "回撤触线时，谁来扣扳机？", "最大回撤预算往往只是写在风控手册里的红线，没有机制保证它在盘中被真正执行。", EMER, EMER_T)]
for i, (k, q, d, c, tint) in enumerate(qs):
    x = Inches(0.6 + i * 4.18)
    box(s, x, Inches(2.15), Inches(3.95), Inches(3.7), fill=WHITE, line=BORDER, round_=True)
    box(s, x, Inches(2.15), Inches(3.95), Inches(0.09), fill=c)
    num_badge(s, x + Inches(0.22), Inches(2.42), i + 1, c)
    txt(s, x + Inches(0.9), Inches(2.5), Inches(2.9), Inches(0.3), [[(k, 9.5, FAINT, False)]])
    txt(s, x + Inches(0.22), Inches(3.1), Inches(3.55), Inches(0.75), [[(q, 14, INK, True)]])
    txt(s, x + Inches(0.22), Inches(3.95), Inches(3.55), Inches(1.75), [[(d, 10, GRAY, False)]], spacing=1.15)
insight(s, "信号「准不准」已被讨论很多，信号「配多少」才是直接决定损益的环节", "教仓位决策 = 教风险预算官", BLUE, BLUE_T)
footer(s, 4, BLUE)

# ================= S5 为什么必须虚拟仿真 =================
s = slide()
box(s, 0, 0, SW, SH, fill=WHITE)
header(s, "01 · 背景与意义", "为什么必须虚拟仿真：真实行情沙盘 × 透明信号注入", BLUE,
       sub="培训规范要求「取材现实、数据有出处」——行情真实可核验，信号透明可复现，二者结合才构成完整沙盘")
card(s, Inches(0.6), Inches(2.1), Inches(5.95), Inches(3.35), BLUE,
     "真实行情（出处明确）",
     "28 只沪深大盘蓝筹 · 410 个交易日（2024-01-02 至 2025-10-28）前复权收盘价\n来源：腾讯财经公开行情接口，任何学生可随时重新下载核验\n覆盖金融/消费/能源/制造/科技，含完整牛熊波动结构",
     head_size=14, body_size=10.5)
card(s, Inches(6.8), Inches(2.1), Inches(5.95), Inches(3.35), EMER,
     "透明信号（种子固定）",
     "同一行情上蒙特卡洛注入 6910 条买/卖信号\n方向命中率 55% 假设（实测 53.95%），种子 20261029\n置信度与真实超额幅度正相关 + 噪声；生成规则完全公开，复跑逐字节一致",
     head_size=14, body_size=10.5)
box(s, Inches(0.6), Inches(5.6), Inches(12.15), Inches(0.42), fill=WHITE, line=BORDER, round_=True)
txt(s, Inches(0.9), Inches(5.66), Inches(11.6), Inches(0.3),
    [[("教学定位：训练「仓位决策」能力而非复刻某段行情——真实行情保证环境复杂性，透明信号保证实验可控性。", 10.5, INK, True)]])
insight(s, "数据有出处是规范性得分点，信号有种子是可复现得分点", "真实为体 · 透明为用", BLUE, BLUE_T, y=Inches(6.15))
footer(s, 5, BLUE)

# ================= S6 章节页 02 =================
divider("02", "EXPERIMENTAL DESIGN", "实验设计",
        ["冻结同一组 AI 信号，只变动「信号→仓位」的转化规则",
         "三组规则对照：满仓等权 / 置信度分档 / 1/4 凯利+回撤干预",
         "统一约束：5bp 单边成本 · 20% 个股上限（硬）· 6% 回撤软触线"],
        EMER, EMER_T)

# ================= S7 实验环境总览（网页截图） =================
s = slide()
box(s, 0, 0, SW, SH, fill=WHITE)
header(s, "02 · 实验设计", "实验环境：平台在线 Notebook，打开即跑", EMER,
       sub="案例正文、数据集、代码、运行结果一体化，学生零安装配置")
pic = s.shapes.add_picture("截图_正文_顶部摘要.png", Inches(0.6), Inches(1.95), height=Inches(3.95))
box(s, Inches(0.6), Inches(1.95), pic.width, Inches(3.95), line=BORDER, line_w=1.2)
txt(s, Inches(0.6), Inches(5.95), Inches(6), Inches(0.3), [[("平台案例正文实拍 · 在线 Notebook", 9, FAINT, False)]])
feats = [("一体化载体", "实验说明、数据文件、代码单元、运行结果与报告要求全部集成在一个 Notebook", EMER),
         ("云端运行", "学生登录即获独立实验环境，预置数据分析工具与依赖，减少本地安装干扰", BLUE),
         ("一键复现", "全部表格与图形由代码单元真实运行生成，无任何手工录入数字", VIOLET)]
for i, (h, b, c) in enumerate(feats):
    y = Inches(2.0 + i * 1.32)
    box(s, Inches(7.15), y, Inches(5.55), Inches(1.15), fill=WHITE, line=BORDER, round_=True)
    box(s, Inches(7.15), y, Inches(0.09), Inches(1.15), fill=c)
    txt(s, Inches(7.45), y + Inches(0.13), Inches(5.1), Inches(0.35), [[(h, 12.5, INK, True)]])
    txt(s, Inches(7.45), y + Inches(0.5), Inches(5.1), Inches(0.6), [[(b, 9.8, GRAY, False)]], spacing=1.1)
insight(s, "评审最关注「方案可比较、过程可追溯」——每个数字都能一键重跑", "环境即教具", EMER, EMER_T)
footer(s, 7, EMER)

# ================= S8 实验设计表 =================
s = slide()
box(s, 0, 0, SW, SH, fill=WHITE)
header(s, "02 · 实验设计", "受控对照实验：信号冻结，只变仓位规则", EMER,
       sub="信号相同，则任何绩效差异都只能归因于仓位规则——干净识别因果关系的关键")
groups = [("A 对照组", "满仓等权（无风险预算）", "有信号即 100% 等权持有", "「All in」的直觉做法", ROSE, ROSE_T),
          ("B 实验组", "置信度分档固定分数", "≥0.7 配 3 成；0.55–0.7 配 2 成", "把置信度直接映射为仓位", BLUE, BLUE_T),
          ("C 实验组", "1/4 凯利 + 回撤干预", "样本内估凯利分数×1/4；触 -6% 线仓位减半", "定量风险预算 + 主动干预", EMER, EMER_T)]
for i, (g, t, d, role, c, tint) in enumerate(groups):
    x = Inches(0.6 + i * 4.18)
    box(s, x, Inches(2.05), Inches(3.95), Inches(3.3), fill=WHITE, line=BORDER, round_=True)
    box(s, x, Inches(2.05), Inches(3.95), Inches(0.09), fill=c)
    txt(s, x + Inches(0.22), Inches(2.3), Inches(3.5), Inches(0.3), [[(g, 10, FAINT, False)]])
    txt(s, x + Inches(0.22), Inches(2.62), Inches(3.5), Inches(0.55), [[(t, 14, INK, True)]])
    txt(s, x + Inches(0.22), Inches(3.3), Inches(3.55), Inches(0.9), [[(d, 10, GRAY, False)]], spacing=1.12)
    box(s, x + Inches(0.22), Inches(4.6), Inches(3.5), Inches(0.55), fill=tint, round_=True)
    txt(s, x + Inches(0.4), Inches(4.68), Inches(3.2), Inches(0.4), [[("角色：" + role, 10, c if c != ROSE else ROSE, True)]])
box(s, Inches(0.6), Inches(5.55), Inches(12.15), Inches(0.42), fill=WHITE, line=BORDER, round_=True)
txt(s, Inches(0.9), Inches(5.61), Inches(11.6), Inches(0.3),
    [[("三组共用硬约束：", 10.5, EMER, True),
      ("硬约束：单票上限 20% · 总仓上限 100% · 单边成本 5bp ｜ 软触线：回撤 -6% 触发减半（不保证不突破）", 10.5, INK, False)]])
insight(s, "任何参数不得使用样本外信息：前 70% 估计，后 30% 检验", "机制的可解释性优先于结果的漂亮", EMER, EMER_T)
footer(s, 8, EMER)

# ================= S9 数据诊断 =================
s = slide()
box(s, 0, 0, SW, SH, fill=WHITE)
header(s, "02 · 实验设计", "信号诊断：置信度存在系统性过度自信", EMER,
       sub="置信度均值 0.871，但各桶真实命中率 50.9%–58.3% 且毫无单调性——这是实验组设计的直接依据")
styled_table(s, Inches(0.6), Inches(2.15), Inches(7.3),
             ["置信度桶", "信号数", "方向命中率"],
             [["[0.5, 0.6)", "446", "58.3%"],
              ["[0.6, 0.7)", "709", "50.9%"],
              ["[0.7, 0.8)", "746", "54.3%"],
              ["[0.8, 0.9)", "712", "54.1%"],
              ["[0.9, 1.0)", "4214（占61%）", "54.0%"]],
             col_w=[2.9, 2.4, 2.0], font_size=11, row_h=0.42, head_fill=EMER)
card(s, Inches(8.25), Inches(2.15), Inches(4.5), Inches(2.45), ROSE,
     "两个层次的同时失真",
     "① 绝对失真：均值 0.871 vs 真实命中 54%——整体过度自信\n② 排序失灵：命中率最高的反而是 [0.5,0.6) 桶——置信度分档没有区分度",
     head_size=13, body_size=10.5)
box(s, Inches(8.25), Inches(4.8), Inches(4.5), Inches(1.3), fill=EMER_T, round_=True)
txt(s, Inches(8.5), Inches(4.95), Inches(4.0), Inches(1.0),
    [[("结论：", 11, EMER, True),
      ("置信度只能作参考输入，须经凯利打折与硬约束双重衰减——实验组 B 将验证这一点。", 11, GRAY, False)]], spacing=1.12)
insight(s, "AI 的置信度不是仓位权重：未经验证的自信会放大为敞口", "先校准，再分配", EMER, EMER_T)
footer(s, 9, EMER)

# ================= S10 章节页 03 =================
divider("03", "FINDINGS & EVIDENCE", "结果与发现",
        ["满仓组的成本之税：年化拖累 8.8%，风险预算组仅 4.3%",
         "样本外反转揭示保险的定价：看得见的机会成本，看不见的防御价值",
         "反事实检验：同时报告「移除干预」的不利情形，拒绝选择性报告"],
        ORANGE, ORAN_T)

# ================= S11 全样本对比 =================
s = slide()
box(s, 0, 0, SW, SH, fill=WHITE)
header(s, "03 · 结果与发现", "全样本对照：同样的信号，完全不同的持有体验", ORANGE,
       sub="410 个交易日 · 三组共用同一信号与硬约束 · 绩效差异全部归因于仓位规则")
s.shapes.add_picture("ppt_fig1_nav.png", Inches(0.6), Inches(2.0), width=Inches(8.1))
kvs = [("23.2% → 18.6%", "满仓组 vs 风险预算组 累计收益：约 4.6 个百分点的「保费」", BLUE),
       ("-12.4% → -9.3%", "最大回撤被压缩 3.1 个百分点，卡玛 1.11 → 1.19", EMER),
       ("0.70 → 0.34", "日均单边换手腰斩——省下的成本就是超额收益", VIOLET)]
for i, (k, v, c) in enumerate(kvs):
    y = Inches(2.1 + i * 1.32)
    box(s, Inches(8.95), y, Inches(3.85), Inches(1.15), fill=WHITE, line=BORDER, round_=True)
    box(s, Inches(8.95), y, Inches(0.09), Inches(1.15), fill=c)
    txt(s, Inches(9.25), y + Inches(0.1), Inches(3.4), Inches(0.4), [[(k, 15, c, True)]])
    txt(s, Inches(9.25), y + Inches(0.52), Inches(3.45), Inches(0.6), [[(v, 9.5, GRAY, False)]], spacing=1.1)
insight(s, "集中不等于低风险：凯利组波动未降，风险预算的价值要用回撤与成本衡量", "怎么买，比买什么更影响体验", ORANGE, ORAN_T)
footer(s, 11, ORANGE)

# ================= S12 成本与回撤 =================
s = slide()
box(s, 0, 0, SW, SH, fill=WHITE)
header(s, "03 · 结果与发现", "交易成本是仓位规则的隐形税", ORANGE,
       sub="5bp 看似微小，但满仓组日均换手 0.70——年化成本拖累 8.8 个百分点，足以吞掉微弱的预测力")
s.shapes.add_picture("ppt_fig2_dd.png", Inches(0.6), Inches(2.0), width=Inches(8.1))
card(s, Inches(8.95), Inches(2.05), Inches(3.85), Inches(3.85), ORANGE,
     "教学要点",
     "· 满仓组年化成本拖累 8.8%，风险预算组仅 4.3%\n· 样本外上行段换手升至 0.91，拖累 11.4%——高频调仓的成本随行情放大\n· 任何「预测很强」的信号，必须先过成本这一关才能变成收益\n· 不做成本扣除的回测没有讨论价值（典型错误引导之一）",
     head_size=13, body_size=9.8)
insight(s, "方向命中率 54% 的信号，在成本面前没有犯错的余地", "先算成本，再谈预测", ORANGE, ORAN_T)
footer(s, 12, ORANGE)

# ================= S13 样本外 =================
s = slide()
box(s, 0, 0, SW, SH, fill=WHITE)
header(s, "03 · 结果与发现", "样本外检验：风险预算是「有价格的保险」", ORANGE,
       sub="后 30% 样本恰为蓝筹上行段——满仓组年化 41.3% vs 风险预算组 13.5%，但回撤 -4.1% vs -2.5%")
pair = [("满仓等权（无预算）", "年化 41.3%", "夏普 2.71", "最大回撤 -4.1%", "日均换手 0.91 · 成本拖累 11.4%", ROSE, ROSE_T),
        ("1/4 凯利 + 回撤干预", "年化 13.5%", "夏普 1.56", "最大回撤 -2.5%", "日均换手 0.38 · 成本拖累 4.8%", EMER, EMER_T)]
for i, (g, a, b, c, d, col, tint) in enumerate(pair):
    x = Inches(0.6 + i * 6.25)
    box(s, x, Inches(2.05), Inches(5.95), Inches(3.1), fill=WHITE, line=BORDER, round_=True)
    box(s, x, Inches(2.05), Inches(5.95), Inches(0.09), fill=col)
    txt(s, x + Inches(0.25), Inches(2.28), Inches(5.4), Inches(0.35), [[(g, 12.5, INK, True)]])
    txt(s, x + Inches(0.25), Inches(2.72), Inches(5.4), Inches(0.6), [[(a, 25, col, True)]])
    txt(s, x + Inches(0.25), Inches(3.5), Inches(5.4), Inches(0.4), [[(b + "    " + c, 12.5, GRAY, True)]])
    txt(s, x + Inches(0.25), Inches(4.05), Inches(5.4), Inches(0.8), [[(d, 10, FAINT, False)]])
box(s, Inches(0.6), Inches(5.35), Inches(12.15), Inches(0.62), fill=ORAN_T, round_=True)
txt(s, Inches(0.9), Inches(5.44), Inches(11.6), Inches(0.5),
    [[("这不是实验失败：", 11, ORANGE, True),
      ("保险的机会成本在平安时期「看得见」，防御价值在回撤时期才兑现——只看一个区间下结论，正是本案例要治的病。", 11, GRAY, False)]], spacing=1.1)
insight(s, "样本外子区间 C 组落后 ≠ 风险预算失败：用保险定价的逻辑辩护或批评", "用两个区间说话，而不是一个", ORANGE, ORAN_T)
footer(s, 13, ORANGE)

# ================= S14 敏感性与反事实 =================
s = slide()
box(s, 0, 0, SW, SH, fill=WHITE)
header(s, "03 · 结果与发现", "敏感性与反事实：干预的价值随仓位规模显现", ORANGE,
       sub="凯利缩放 1/8 → 满仓 × 干预有无（无干预=反事实）：同一策略、同一数据，差异全部归因于干预机制")
s.shapes.add_picture("ppt_fig3_sens.png", Inches(0.6), Inches(2.0), width=Inches(8.1))
card(s, Inches(8.95), Inches(2.05), Inches(3.85), Inches(3.85), VIOLET,
     "三个梯度结论",
     "· 1/8 凯利：仓位本轻，干预与否完全相同\n· 1/2 凯利：干预把回撤从 -15.9% 压到 -12.3%\n· 满仓凯利：无干预回撤 -15.7%，干预压回 -12.6%\n· 但干预不是免费的：1/4 档收益 22.7%→18.6% 而回撤未变——预算线对轻仓位太松",
     head_size=13, body_size=9.8)
insight(s, "仓位规模 × 干预线要成对设计：只调一个是新手做法", "保费要与敞口匹配", ORANGE, ORAN_T)
footer(s, 14, ORANGE)

# ================= S15 干预线设计 =================
s = slide()
box(s, 0, 0, SW, SH, fill=WHITE)
header(s, "03 · 结果与发现", "回撤预算线设在哪里？—— 干预线敏感性", ORANGE,
       sub="固定 1/4 凯利，只移动预算线：-3% / -6% / -9%（全样本）")
lines_ = [("-3% 预算线", "卡玛比率 1.49", "回撤 -7.7%", "干预频繁、换手更高，把回撤真正压住了", EMER, EMER_T),
          ("-6% 预算线", "卡玛比率 1.19", "回撤 -9.3%", "触线即减半的软干预未能把回撤压进预算——留作改进作业", BLUE, BLUE_T),
          ("-9% 预算线", "卡玛比率 1.41", "回撤 -9.3%", "触发太少，干预机制名存实亡——太松等于没有", FAINT, TINT)]
for i, (t, k, dd, d, col, tint) in enumerate(lines_):
    x = Inches(0.6 + i * 4.18)
    box(s, x, Inches(2.05), Inches(3.95), Inches(3.55), fill=WHITE, line=BORDER, round_=True)
    box(s, x, Inches(2.05), Inches(3.95), Inches(0.09), fill=col)
    txt(s, x + Inches(0.22), Inches(2.28), Inches(3.5), Inches(0.4), [[(t, 13.5, INK, True)]])
    txt(s, x + Inches(0.22), Inches(2.75), Inches(3.5), Inches(0.55), [[(k, 19, col, True)]])
    txt(s, x + Inches(0.22), Inches(3.38), Inches(3.5), Inches(0.35), [[(dd, 11, GRAY, True)]])
    txt(s, x + Inches(0.22), Inches(3.82), Inches(3.55), Inches(1.6), [[(d, 9.8, GRAY, False)]], spacing=1.14)
insight(s, "干预线不是越紧越好：正确的问题是「线的位置与仓位规模是否匹配」", "成对设计，协同调参", ORANGE, ORAN_T)
footer(s, 15, ORANGE)

# ================= S16 章节页 04 =================
divider("04", "TEACHING & ASSESSMENT", "教学组织与评价",
        ["分层任务：基础必做 / 进阶选做 / 挑战任务，结果不唯一",
         "典型错误引导：前视偏差 · 忽略成本 · 样本内偷看 · 选择性报告",
         "四维评价量规：过程可追溯 · 解释合理性 · 实验完整性 · 报告规范"],
        VIOLET, VIOL_T)

# ================= S17 分层任务 =================
s = slide()
box(s, 0, 0, SW, SH, fill=WHITE)
header(s, "04 · 教学组织与评价", "分层任务设计：每个学生都有可到达的「最近发展区」", VIOLET,
       sub="必做保底、选做分层、挑战拔尖——实验结果不唯一，干预规则的设计空间向学生开放")
tiers = [("基础任务（必做）", EMER, EMER_T,
          "· 改进干预机制：替换「触线减半」，重跑全部实验并用数据论证\n· 完整复现：从数据生成到结论一键重跑，验证与案例报告一致"),
         ("进阶任务（选做）", BLUE, BLUE_T,
          "· 把信号命中率降到 50%（无预测力）重做实验，解释为何风险预算组相对抗跌\n· 以卡玛比率为目标在（凯利缩放×预算线）网格上搜索，讨论是否过拟合"),
         ("挑战任务（拔尖）", VIOLET, VIOL_T,
          "· 接入真实公开信号源（或自研大模型信号）在相同规则下重跑，分解「信号质量」与「仓位规则」各自的绩效贡献")]
for i, (t, col, tint, b) in enumerate(tiers):
    x = Inches(0.6 + i * 4.18)
    box(s, x, Inches(2.05), Inches(3.95), Inches(3.85), fill=WHITE, line=BORDER, round_=True)
    box(s, x, Inches(2.05), Inches(3.95), Inches(0.48), fill=col)
    txt(s, x, Inches(2.12), Inches(3.95), Inches(0.36), [[(t, 12.5, WHITE, True)]], align=PP_ALIGN.CENTER)
    txt(s, x + Inches(0.22), Inches(2.72), Inches(3.55), Inches(3.0),
        [[(ln, 10, GRAY, False)] for ln in b.split("\n")], spacing=1.2)
insight(s, "开放的干预规则设计空间 = 「结果不唯一」的评审要求", "人人有得做，拔尖有挑战", VIOLET, VIOL_T)
footer(s, 17, VIOLET)

# ================= S18 典型错误引导 =================
s = slide()
box(s, 0, 0, SW, SH, fill=WHITE)
header(s, "04 · 教学组织与评价", "典型错误引导：实验前必读的四个坑", VIOLET,
       sub="每个坑都给出「错误现象 → 为什么错 → 如何自查」的三段式排查路径")
errs = [("前视偏差", "信号在 t 日收盘后生成，只能作用于 t+1 及以后收益；直接对齐 t 日收益会让所有策略虚高。", "检查信号生效时间是否对齐到「收盘后」"),
        ("忽略成本", "5bp 很小，但满仓组年化拖累 8.8 个百分点——不做成本扣除的回测没有讨论价值。", "核对回测引擎是否在换仓日扣除换手×5bp"),
        ("样本内偷看", "用全样本估计凯利分数等于偷看答案，样本外检验形同虚设。", "确认所有参数只使用前 70% 样本"),
        ("选择性报告", "只展示「干预有效」的情形是评审最反感的做法。", "同时报告反事实（移除干预）结果并讨论")]
for i, (t, why, check) in enumerate(errs):
    x = Inches(0.6 + (i % 2) * 6.25); y = Inches(2.05 + (i // 2) * 1.95)
    box(s, x, y, Inches(5.95), Inches(1.8), fill=WHITE, line=BORDER, round_=True)
    box(s, x, y, Inches(0.09), Inches(1.8), fill=ROSE)
    txt(s, x + Inches(0.28), y + Inches(0.13), Inches(5.4), Inches(0.35), [[(f"坑 {i+1}：{t}", 12.5, INK, True)]])
    txt(s, x + Inches(0.28), y + Inches(0.52), Inches(5.45), Inches(0.8), [[(why, 9.5, GRAY, False)]], spacing=1.1)
    txt(s, x + Inches(0.28), y + Inches(1.36), Inches(5.45), Inches(0.35),
        [[("自查：", 9.5, EMER, True), (check, 9.5, GRAY, False)]])
insight(s, "把坑写进指导手册，是把「翻车点」变成「教学点」", "错误是最好的教具", VIOLET, VIOL_T)
footer(s, 18, VIOLET)

# ================= S19 评价量规 =================
s = slide()
box(s, 0, 0, SW, SH, fill=WHITE)
header(s, "04 · 教学组织与评价", "四维评价量规：重过程、重解释，不重收益率", VIOLET,
       sub="评分侧重「过程可追溯」与「解释合理性」——高收益不加分，不可复现扣大分")
styled_table(s, Inches(0.6), Inches(2.05), Inches(12.15),
             ["维度", "权重", "优（90–100）", "合格（60–75）", "不合格（<60）"],
             [["过程可追溯", "30%", "每个数字可由代码一键复现，随机种子与参数全记录", "主体可复现，个别参数缺失", "结果无法重现或依赖手工录入"],
              ["解释合理性", "30%", "结论与证据严格匹配，主动讨论反事实与局限", "结论基本有据，局限讨论薄弱", "结论脱离数据或回避不利证据"],
              ["实验完整性", "20%", "对照+样本外+敏感性+反事实四件套齐全", "缺一项但有替代论证", "仅单一区间或单一参数下结论"],
              ["报告规范", "20%", "结构完整、图表清晰、匿名合规、口径一致", "结构完整但图表粗糙", "缺环节或出现身份信息"]],
             col_w=[1.9, 1.0, 3.6, 3.2, 2.45], font_size=9.5, row_h=0.6, head_fill=VIOLET)
insight(s, "量规先行发布：学生知道怎么被评，才知道怎么学", "评过程，不评运气", VIOLET, VIOL_T)
footer(s, 19, VIOLET)

# ================= S20 结语 =================
s = slide()
box(s, 0, 0, SW, SH, fill=INK)
oval(s, Inches(9.8), Inches(-2.4), Inches(6.0), RGBColor(0x1E, 0x29, 0x3B))
oval(s, Inches(11.3), Inches(-1.2), Inches(3.0), RGBColor(0x25, 0x63, 0xEB))
box(s, Inches(0.9), Inches(1.0), Inches(0.9), Inches(0.09), fill=EMER)
txt(s, Inches(0.9), Inches(1.3), Inches(5), Inches(0.35), [[("CLOSING", 12, RGBColor(0x94, 0xA3, 0xB8), True)]])
txt(s, Inches(0.9), Inches(1.65), Inches(5), Inches(0.7), [[("结语", 30, WHITE, True)]])
txt(s, Inches(0.9), Inches(2.9), Inches(11.6), Inches(1.7),
    [[("让每一名学生在真金白银之前，先在虚拟仿真沙盘上亲历一次", 19, RGBColor(0xCB, 0xD5, 0xE1), False)],
     [("「预测力 ≠ 盈利能力」的教训，带走一套可验证、可复核、可追溯的决策纪律。", 19, WHITE, True)]], spacing=1.45)
box(s, Inches(0.9), Inches(4.85), Inches(11.5), Inches(0.02), fill=RGBColor(0x33, 0x41, 0x55))
txt(s, Inches(0.9), Inches(5.1), Inches(11.6), Inches(0.5),
    [[("AI 是外骨骼，扣动扳机的必须是懂风险预算的人。", 15, EMER, True)]])
txt(s, Inches(0.9), Inches(6.15), Inches(11.5), Inches(0.35),
    [[("感谢各位评委老师的聆听与指正", 12, RGBColor(0x94, 0xA3, 0xB8), False)]])
txt(s, Inches(0.9), Inches(6.55), Inches(11.5), Inches(0.35),
    [[("AI 选股信号该配多大仓位？——虚拟仿真下的风险预算干预实验", 10.5, RGBColor(0x64, 0x74, 0x8B), False)]])
box(s, 0, Inches(7.4), SW, Inches(0.1), fill=EMER)
txt(s, Inches(11.6), Inches(7.0), Inches(1.2), Inches(0.35), [[(f"20 / {TOTAL}", 10, EMER, True)]], align=PP_ALIGN.RIGHT)

prs.save("教学指导配套PPT_v2.pptx")
print("saved 教学指导配套PPT_v2.pptx | slides:", len(prs.slides._sldIdLst))
