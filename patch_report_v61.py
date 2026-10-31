# -*- coding: utf-8 -*-
"""案例报告 v6.1 梳理补丁：动机收束 + 方法工具箱 + 市场应用与产品化 + 可提高清单。只增不减。"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

SRC = "/Users/Admin/Documents/kimi/tasks/2026-10-24/14-58-19-b5e2402f/提交暂存/案例报告.docx"
d = Document(SRC)

def east(el, font):
    rPr = el.get_or_add_rPr(); rPr.get_or_add_rFonts().set(qn("w:eastAsia"), font)

def find_para(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)

def heading_before(anchor, text, size=15):
    p = anchor.insert_paragraph_before()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    r.font.name = "Times New Roman"; east(r._element, "黑体")
    return p

def body_before(anchor, text, size=13, indent=True, bold=False):
    p = anchor.insert_paragraph_before()
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold
    r.font.name = "Times New Roman"; east(r._element, "宋体")
    return p

def table_before(anchor, rows, size=9):
    ncol = max(len(r) for r in rows)
    tb = d.add_table(rows=len(rows), cols=ncol)
    try: tb.style = d.tables[0].style
    except Exception: tb.style = "Table Grid"
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = tb.cell(i, j); cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(row[j] if j < len(row) else "")
            r.font.size = Pt(size); r.font.name = "Times New Roman"; east(r._element, "宋体")
            if i == 0: r.bold = True
    anchor._p.addprevious(tb._tbl)
    return tb

def caption_before(anchor, text):
    p = anchor.insert_paragraph_before()
    p.alignment = 1
    r = p.add_run(text); r.bold = True; r.font.size = Pt(10.5)
    r.font.name = "Times New Roman"; east(r._element, "黑体")

# ============ ① 为什么做：1.1 末尾研究动机收束段 ============
a12 = find_para("1.2　研究问题")
body_before(a12, "本研究要做这件事，动机有三。其一，技术现状与教学空白并存：大模型投研工作台已批量产出\u201c买什么\u201d的信号，但\u201c买多少\u201d的转化环节既无训练环境、也无考核标准，课程体系里几乎完全缺席。其二，投资者行为的真实痛点：相当比例的个人投资者以满仓方式应对一切行情——行情好时把贝塔误认为能力，行情差时回撤失控；这不是知识缺口，而是\u201c从信号到资金\u201d的决策纪律缺口。其三，AI 治理的时代命题：大模型口诉置信度被证实存在系统性过度自信，若直接以置信度驱动仓位，等于把模型的自信误当作模型的能力。把这三件事合在同一间实验室里，就是本研究的出发点：用一个受控、可复现、可证伪的虚拟仿真沙盘，把\u201c信号给出之后买多少\u201d教成一门可训练、可考核的手艺。", indent=True)

# ============ ② 理论依据与工具方法：2.6 方法工具箱 ============
a3 = find_para("3　数据与研究设计")
heading_before(a3, "2.6　方法工具箱：本研究用什么理论依据与工具方法")
body_before(a3, "第 2.1\u20132.5 节回答\u201c站在哪条学术河流里\u201d，本节回答\u201c用什么工具干活\u201d。本研究的方法论由一张工具箱给出：每一件工具都对应一个必须回答的问题，都有文献血统（见参考文献），也都在正文中留下了可核验的输出。")
caption_before(a3, "表 12　方法工具箱：每件工具回答一个问题")
table_before(a3, [
    ["工具方法", "回答的问题", "理论依据", "正文落点"],
    ["受控对照实验", "绩效差异到底由什么造成？", "实验经济学：冻结信号、只变仓位规则，识别干净因果", "3.3、4.2"],
    ["样本外切分", "结论是不是只对样本内成立？", "回测纪律：后 30% 严格 OOS 验证", "4.3"],
    ["反事实检验", "干预到底值多少钱？", "因果推断：移除干预重跑同一引擎", "4.4"],
    ["参数敏感性网格", "结论依赖特定参数吗？", "稳健性检验：凯利缩放 × 干预有无全网格", "4.4"],
    ["块状自助法", "差异是运气还是机制？", "统计诚实：1000 次重抽样给出 95% CI", "5.5"],
    ["参数预登记", "网格最优可信吗？", "多重检验修正：网格最优必须样本外复验", "1.4、5.5"],
    ["成本扫描与凸性冲击", "结论在真实成本结构下还成立吗？", "市场微结构：平方根律情景", "5.1、5.6"],
    ["敞口匹配基准", "改善来自机制还是仓位更轻？", "基准构造：主动自我证伪", "5.4"],
    ["置信度校准与再校准", "AI 的自信能直接用吗？", "校准理论：分桶诊断 + 保序再校准", "4.1、5.2"],
    ["锚点自检", "每个数字可靠吗？", "复现工程：24/24 锚点一键核验", "附录 A"],
])
body_before(a3, "工具箱的使用纪律只有一条：**先问这个问题有没有被回答的证据，再谈结论**。凡是工具箱覆盖不到的主张（例如\u201c实盘一定能赚钱\u201d），正文一律不做断言——这份克制本身就是方法的一部分。")

# ============ ④⑤ 市场应用与产品化：6.5 ============
aref = find_para("参考文献")
heading_before(aref, "6.5　市场应用场景与产品化路径")
body_before(aref, "本研究的产出不止于一份教学案例：\u201c信号→校准→仓位→复核\u201d的完整工作流，配上可一键复现的工程内核（单一事实源 run_all.py、24/24 数字锚点、15 组结果 CSV），本身就是一个可交付产品的雏形。其在市场中的落地场景如表 13。")
caption_before(aref, "表 13　应用场景与产品形态")
table_before(aref, [
    ["场景", "谁在用", "用的是什么", "形态"],
    ["高校实验教学", "投资学/量化/金融科技课程", "本案例全套：沙盘+三组对照+参数实验室+四维量规", "课程包（已交付）"],
    ["券商与机构投教", "投顾培训、客户风险教育", "开篇七问 + 画像问卷 + 回撤预算线演示", "工作坊/内训模块"],
    ["个人投资者工具", "以两年期工资型资金为主的散户", "五维画像自评 → 画像驱动仓位规则 → 复利与回撤可视化", "APP/小程序（规划中）"],
    ["机构研报服务", "机构与专业投资者", "覆盖池 300 只、日频\u201c信号→校准→仓位\u201d研报生产线", "订阅制（中期）"],
    ["金融数字孪生", "财富管理与智能投顾平台", "画像 × 行为校验驱动的个性化仓位引擎", "SaaS 接口（远期）"],
])
body_before(aref, "盈利路径与阶段一一对应：近期以院校课程包授权与机构内训为主，边际成本近乎为零——全部材料已按\u201c他校教师可开课\u201d标准打磨；中期以日频研报订阅为现金流，其质量由本研究的校准工作流背书：每份研报附带置信度校准状态与仓位建议的来源可追溯性，这正是与市面\u201c黑箱荐股\u201d的根本区别；远期的金融数字孪生 SaaS 把画像引擎以接口形式输出给财富管理平台，按调用量计费。需要强调的是：**产品化的每一步都建立在本研究的公共工程之上**——随机种子、参数预登记、锚点自检与全部结果 CSV 均已公开，可复现性既是学术纪律，也是产品的第一卖点和护城河。")
body_before(aref, "最后，如实列出本工具当前的可提高之处，与 6.3 的局限一一衔接：一是信号层仍是透明模拟，接入真实大模型在线信号后，校准工作流需要从离线批处理升级为在线闭环；二是画像引擎只有问卷自评单通道，行为数据校验通道尚未建立，画像分辨率有限；三是执行层仍以参数化成本刻画，未接入实时盘口；四是教学证据目前以设计质量与可复现性为主，真实学习者的大样本前后测数据有待积累。这四项分别对应 6.4 的方向二、一、三与教学实践——**工具是进化的，路线是未完待续的**。")

d.save(SRC)
print("✓ v6.1 梳理补丁完成")

# 校验
d2 = Document(SRC)
full = "\n".join(p.text for p in d2.paragraphs)
for a in ["2.6　方法工具箱", "6.5　市场应用场景与产品化路径", "表 12", "表 13", "动机有三", "可提高之处", "护城河"]:
    print(("✓" if a in full else "✗"), a)
print("段落:", len(d2.paragraphs), "表格:", len(d2.tables))
