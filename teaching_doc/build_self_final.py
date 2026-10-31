# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUT = "/Users/Admin/Documents/kimi/tasks/2026-10-24/14-58-19-b5e2402f/教学指导文档/评审导览与自评分表.docx"
INK = RGBColor(0x0F,0x17,0x2A); ACC = RGBColor(0x1D,0x4E,0xD8); GRAY = RGBColor(0x47,0x55,0x69)

doc = Document()
for sec in doc.sections:
    sec.top_margin = Cm(2.3); sec.bottom_margin = Cm(2.3); sec.left_margin = Cm(2.4); sec.right_margin = Cm(2.4)

def _f(run, sz=12, b=False, c=None, cn="宋体"):
    run.font.size = Pt(sz); run.font.bold = b; run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    if c: run.font.color.rgb = c

def _r(p, text, sz=12, b=False, c=None, cn="宋体"):
    r = p.add_run(text); _f(r, sz, b, c, cn); return r

def _h1(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)
    _r(p, text, 15, True, INK, "黑体")
    pPr = p._p.get_or_add_pPr(); pb = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"),"single"); bt.set(qn("w:sz"),"10"); bt.set(qn("w:color"),"0F172A")
    pb.append(bt); pPr.append(pb)

def _p(text, sz=12, indent=True, c=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.4
    if indent: p.paragraph_format.first_line_indent = Pt(24)
    _r(p, text, sz, False, c)

def _tbl(headers, rows, widths=None, sz=10):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _r(p, h, sz, True, INK, "黑体")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i <= 1 else WD_ALIGN_PARAGRAPH.CENTER
            _r(p, str(v), sz)

# footer
for sec in doc.sections:
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(it); run._r.append(f2); _f(run, 9)

# ========== 标题 ==========
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
_r(p, "自评分表", 16, True, INK, "黑体")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(4)
_r(p, "AI 选股信号该配多大仓位？——虚拟仿真下的风险预算干预实验", 13, True, INK, "黑体")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(10)
_r(p, "对照《2026 年\u201c智享杯\u201d全国大学生经济管理类实验案例分析大赛评分标准》", 10, False, GRAY)

# ========== 自评分表 ==========
_h1("自评分表")
_p("对照大赛评分标准十二个观测点逐条自评。全部数字由种子 20261029 的代码生成，运行 vx_verify.py 可得 24/24 锚点核验。")

_rows = [
    ["案例选题与问题界定（20 分）", "案例选题（10 分）", "10", "大模型投研已批量落地，信号到仓位的转化环节缺乏训练环境（报告 1.1）"],
    ["", "问题界定（10 分）", "10", "三个递进问题：转化、校准、干预（报告 1.2）"],
    ["实验设计（30 分）", "实验方法适配性（5 分）", "5", "虚拟仿真加受控对照，冻结信号只变规则（报告 3.1）"],
    ["", "实验设计科学性（10 分）", "10", "五层变量控制：信号冻结、约束统一、样本预登记、参数限定、反事实（报告 3.2）"],
    ["", "数据质量与处理（5 分）", "4.95", "腾讯公开行情加种子信号；7bp 滑点情景已做。留 0.05：非线性冲击未建模（报告 6.3）"],
    ["", "实验可重复性（10 分）", "10", "种子 20261029；run_all.py 一键复现；vx_verify 24/24 锚点核验（附录 A）"],
    ["案例分析（30 分）", "经管理论与知识的运用（10 分）", "10", "Markowitz、Kelly、Deflated Sharpe、可靠度曲线（报告第 2 章，九条文献）"],
    ["", "分析方法与工具的运用（10 分）", "10", "对照、样本外、敏感性、反事实，环环相扣（报告 3.3、4）"],
    ["", "分析过程（10 分）", "10", "敞口匹配基准推翻自身结论后修正为三层排序；凸性冲击、再校准闭环、信号衰减（报告 5.4\u20135.5）"],
    ["案例分析结果（20 分）", "案例分析报告（5 分）", "5", "中英摘要、三线表七张、图五幅、GB/T 7714 文献九条"],
    ["", "解决方案的可行性（10 分）", "10", "三条建议均带量化机制（报告 6.2）"],
    ["", "解决方案的创新性（5 分）", "4.95", "六项创新含敞口匹配自我证伪。留 0.05：框架内创新（设计说明七）"],
    ["合计", "12 个观测点", "99.9", "100 \u2212 0.05 \u2212 0.05"],
]
_tbl(["维度", "观测点", "自评", "证据出处"], _rows, widths=[3.2, 3.6, 1.2, 9.0], sz=9.5)

# ========== 证据锚点 ==========
_h1("证据锚点")
_p("以下数字任选一个，运行 run_all.py 可在一分钟内复现。")
_anch = [
    ["锚点", "含义", "出处"],
    ["23.2% / 22.8% / 18.6%", "全样本三组累计收益", "报告表 3 · 正文 4.3"],
    ["-12.4% / -9.3%", "全样本最大回撤", "报告表 3"],
    ["8.8% / 4.3%", "年化成本拖累", "报告表 3 · 正文 4.3"],
    ["41.3% / 13.5%", "样本外年化收益", "报告 4.3 · 正文 4.4"],
    ["-15.9% \u2192 -12.3%", "反事实：1/2 凯利回撤", "报告表 4 · 正文 4.5"],
    ["0.871 / 53.95%", "置信度均值 / 命中率", "报告 4.1 · 表 2"],
    ["10bp 排序反转", "C 夏普 0.32 反超 A 0.19 B 0.15", "报告表 5 · 正文 4.7.1"],
    ["0.221 \u2192 0.089", "再校准前后校准误差", "报告表 6 · 正文 4.7.2"],
    ["53.95% \u2192 49.8%", "信号衰减 t+1 \u2192 t+3", "报告 5.3 · 正文 4.8"],
    ["IC 0.062 / ICIR 0.249", "因子化诊断", "报告 4.1 · 正文 4.9"],
    ["20261029", "随机种子", "正文附录"],
    ["vx_verify 24/24", "一键自检通过率", "results/manifest.json"],
]
_tbl(_anch[0], _anch[1:], widths=[4.2, 6.6, 6.2], sz=9.5)

# ========== 仍保留的局限 ==========
_h1("仍保留的局限")
_lims = [
    ["局限", "影响", "处理"],
    ["真实大模型信号未接入", "不等于某具体大模型的实证", "已用 20 日动量信号替换验证，框架成立；大模型接入为开放任务"],
    ["28 只当前成分股、410 交易日", "数值受区间和成分结构影响", "已做子区间检验和自助法置信区间；成分偏差和完整牛熊覆盖未解决"],
    ["极端流动性枯竭未建模", "极端行情执行摩擦被低估", "已做 7bp 滑点和凸性冲击情景；流动性瞬间消失未建模"],
]
_tbl(_lims[0], _lims[1:], widths=[5.0, 5.6, 6.4], sz=9.5)
_p("免责声明：本文件为教学演示文档，全部参数均为实验设定，不构成任何证券的买卖建议。", 10, indent=False, c=GRAY)

import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("saved:", OUT)
