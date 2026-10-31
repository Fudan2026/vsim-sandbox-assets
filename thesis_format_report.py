# -*- coding: utf-8 -*-
"""案例报告博士论文排版：封面 + 目录(两遍法真实页码) + 章前换页 + 全表框线 + 页脚页码。匿名封面。"""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PATH = "/Users/Admin/Documents/kimi/tasks/2026-10-24/14-58-19-b5e2402f/提交暂存/案例报告.docx"
d = Document(PATH)

def east(el, font):
    rPr = el.get_or_add_rPr(); rPr.get_or_add_rFonts().set(qn("w:eastAsia"), font)

# ---------- 0) 通用：全表加框线（w:tblBorders 六边） ----------
def set_grid(tbl, sz=6, color="000000"):
    tblPr = tbl._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None: tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        borders.append(e)
    tblPr.append(borders)

for t in d.tables:
    set_grid(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
print("表格框线:", len(d.tables), "张全部加框")

# ---------- 1) 页脚页码（封面除外） ----------
sec = d.sections[0]
sec.different_first_page_header_footer = True
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " PAGE "
f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
run._element.append(f1); run._element.append(it); run._element.append(f2)
run.font.size = Pt(9); run.font.name = "Times New Roman"; east(run._element, "宋体")
print("页脚页码: 已添加（封面首页无页码）")

# ---------- 2) 封面（插入到文档最前） ----------
first = d.paragraphs[0]

def cover_line(anchor, text, size, bold=True, font="黑体", before=0, after=6, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = anchor.insert_paragraph_before()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    r.font.name = "Times New Roman"; east(r._element, font)
    return p

cover_line(first, "第九届“智享杯”全国大学生经济管理类实验案例分析大赛", 16, True, "黑体", before=24, after=4)
cover_line(first, "（学生赛区 · 数据分析案例）", 14, False, "黑体", after=90)
cover_line(first, "AI 选股信号该配多大仓位？", 26, True, "黑体", before=60, after=10)
cover_line(first, "——虚拟仿真下的风险预算干预实验", 20, True, "黑体", after=30)
cover_line(first, "实验教学案例分析报告", 18, True, "仿宋", after=90)
cover_line(first, "参赛团队：大赛匿名评审版", 14, False, "仿宋", after=8)
cover_line(first, "完成日期：2026 年 10 月 31 日", 14, False, "仿宋", after=0)

# ---------- 3) 目录页（占位页码，第二遍回填） ----------
CHAPTERS = ["1　引言", "2　文献与理论基础", "3　数据与研究设计", "4　实证结果", "5　拓展分析",
            "6　结论与启示", "参考文献", "附录 A　数据与实验可复现性声明", "附录 B　符号与参数表"]
# 节级标题自动收集（13pt 粗体、形如 N.N）
sections = []
for p in d.paragraphs:
    t = p.text.strip()
    if not t or len(t) > 46: continue
    bold = any(r.bold for r in p.runs if r.text.strip())
    sz = next((r.font.size.pt for r in p.runs if r.font.size), None)
    if bold and sz == 13.0 and t[:2] in ("1.","2.","3.","4.","5.","6.") and "　" in t:
        sections.append(t)

# 目 录 标题
toc_head = first.insert_paragraph_before()
toc_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
toc_head.paragraph_format.page_break_before = True
toc_head.paragraph_format.space_after = Pt(18)
r = toc_head.add_run("目　　录"); r.bold = True; r.font.size = Pt(22)
r.font.name = "Times New Roman"; east(r._element, "黑体")

def toc_entry(anchor, text, page, level):
    p = anchor.insert_paragraph_before()
    pf = p.paragraph_format
    pf.left_indent = Cm(0 if level == 1 else 0.75)
    pf.space_after = Pt(4)
    pf.tab_stops.add_tab_stop(Cm(15.4), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    r1 = p.add_run(text); r1.font.size = Pt(12 if level == 1 else 11)
    r1.bold = (level == 1); r1.font.name = "Times New Roman"; east(r1._element, "宋体")
    r2 = p.add_run("\t" + str(page)); r2.font.size = Pt(12 if level == 1 else 11)
    r2.font.name = "Times New Roman"; east(r2._element, "宋体")
    return p

# 目录顺序：按文档中出现位置排列章与节（章先收集，节按其文本前缀归属）
toc_plan = []
for p in d.paragraphs:
    t = p.text.strip()
    if t in CHAPTERS: toc_plan.append((1, t))
    elif t in sections: toc_plan.append((2, t))
for lvl, t in toc_plan:
    toc_entry(first, t, 0, lvl)
print(f"目录: {len(toc_plan)} 条（{sum(1 for l,_ in toc_plan if l==1)} 章 + {sum(1 for l,_ in toc_plan if l==2)} 节）")

# ---------- 4) 章前换页 ----------
targets = CHAPTERS + ["Abstract"]
count = 0
for p in d.paragraphs:
    t = p.text.strip()
    if t in CHAPTERS or t == "Abstract":
        p.paragraph_format.page_break_before = True
        count += 1
# 正文标题段（文档正文第一段，位于目录之后）：换页 + Abstract 前换页已在上面
# 找正文标题段 = 目录之后第一段（以 "AI 选股信号该配多大仓位？" 开头且非封面行）
seen_toc = False
for p in d.paragraphs:
    if p.text.strip() == "目　　录": seen_toc = True; continue
    if seen_toc and p.text.strip().startswith("AI 选股信号该配多大仓位"):
        p.paragraph_format.page_break_before = True
        count += 1
        break
print("章前换页:", count, "处")

d.save(PATH)
print("✓ 第一遍完成:", PATH)
